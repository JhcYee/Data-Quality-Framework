"""Excel dashboard: Sheet 1 rule pass/fail + % records affected per rule;
Sheet 2 recommended actions; Sheet 3 column profile, with conditional
formatting.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

from ..anomalies.types import AnomalyResult
from ..expectations import ValidationOutcome, describe_params
from ..profiling import DatasetProfile
from ..recommendations import Recommendation

HEADER_FILL = PatternFill(start_color="1F2933", end_color="1F2933", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def _autofit(ws):
    for column_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 60)


def build_excel_scorecard(
    anomaly_results: list[AnomalyResult],
    validation_outcomes: list[ValidationOutcome],
    recommendations: list[Recommendation],
    profile: DatasetProfile,
) -> bytes:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Rule Summary"
    ws1.append(["Check", "Type", "Result", "% Records Affected", "Details"])
    total_rows = profile.n_rows or 1

    for r in anomaly_results:
        pct = r.affected_row_count / total_rows
        ws1.append([r.check_name, "anomaly", "PASS" if r.passed else "FAIL", pct, r.summary])
    for v in validation_outcomes:
        label = v.expectation_type + (f" ({v.column})" if v.column else "")
        rule = describe_params(v.params)
        details = v.summary + (f" — rule: {rule}" if rule else "")
        ws1.append([label, f"validation rule ({v.source})", "PASS" if v.success else "FAIL", v.unexpected_percent, details])

    _style_header(ws1)
    max_row = ws1.max_row
    if max_row > 1:
        ws1.conditional_formatting.add(
            f"C2:C{max_row}", CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL)
        )
        ws1.conditional_formatting.add(
            f"C2:C{max_row}", CellIsRule(operator="equal", formula=['"PASS"'], fill=PASS_FILL)
        )
    for row in ws1.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = "0.0%"
    _autofit(ws1)

    ws2 = wb.create_sheet("Recommended Actions")
    ws2.append(["Column", "Issue", "Rows affected", "% of rows", "Why it was flagged", "Recommended action"])
    for rec in recommendations:
        ws2.append(
            [rec.column or "(table-level)", rec.issue, rec.affected_rows, rec.pct_of_rows, rec.finding, rec.action]
        )
    _style_header(ws2)
    for row in ws2.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = "0.0%"
    _autofit(ws2)

    ws3 = wb.create_sheet("Column Profile")
    ws3.append(["Column", "Dtype", "Null %", "Unique %"])
    for col, c in profile.columns.items():
        ws3.append([col, c.dtype, c.null_pct, c.unique_pct])
    _style_header(ws3)
    for row in ws3.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = "0.0%"
    _autofit(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
