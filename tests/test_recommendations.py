from dq_framework.anomalies.types import AnomalyResult
from dq_framework.recommendations import build_recommendations
from dq_framework.reporting.markdown_actions import render_markdown_actions


def _fail(check_name: str, n: int, summary: str = "evidence") -> AnomalyResult:
    return AnomalyResult(check_name=check_name, passed=False, summary=summary, affected_row_count=n)


def test_passing_checks_produce_no_recommendations():
    ok = AnomalyResult("nulls:a", True, "0.0% null", 0)
    assert build_recommendations([ok], n_rows=100) == []


def test_each_failing_check_kind_maps_to_an_issue_and_column():
    results = [
        _fail("nulls:age", 11),
        _fail("duplicates:exact_rows", 8),
        _fail("duplicates:key[id]", 16),
        _fail("outliers:score", 3),
        _fail("categorical_standardization:borough", 124),
        _fail("typos:borough", 1),
        _fail("referential:borough_code->code", 38),
        _fail("consistency:close_date>=signup_date", 31),
        _fail("schema_drift", 0),
    ]
    recs = {r.issue: r for r in build_recommendations(results, n_rows=204)}

    assert recs["Missing values"].column == "age"
    assert recs["Duplicate rows"].column is None
    assert recs["Duplicate keys"].column == "id"
    assert recs["Outliers"].column == "score"
    assert recs["Inconsistent casing/whitespace"].column == "borough"
    assert recs["Possible typos"].column == "borough"
    assert recs["Referential integrity break"].column == "borough_code"
    assert recs["Cross-column inconsistency"].column == "close_date, signup_date"
    assert recs["Schema drift"].column is None
    assert all(r.action for r in recs.values())


def test_sorted_by_rows_affected_descending():
    recs = build_recommendations([_fail("nulls:a", 5), _fail("nulls:b", 50), _fail("nulls:c", 20)], n_rows=100)
    assert [r.column for r in recs] == ["b", "c", "a"]


def test_percentage_of_rows_computed():
    (rec,) = build_recommendations([_fail("nulls:a", 25)], n_rows=100)
    assert rec.pct_of_rows == 0.25


def test_evidence_carried_through_from_the_finding():
    (rec,) = build_recommendations([_fail("nulls:a", 3, summary="3.0% null (3 of 100 rows)")], n_rows=100)
    assert rec.finding == "3.0% null (3 of 100 rows)"


def test_coercion_failures_and_placeholders_become_recommendations():
    recs = build_recommendations(
        [], n_rows=100, coercion_failure_counts={"zip": 4}, sentinel_null_counts={"age": 11}
    )
    by_issue = {r.issue: r for r in recs}
    assert by_issue["Values not matching the confirmed type"].affected_rows == 4
    assert by_issue["Placeholder values treated as null"].column == "age"


def test_markdown_lists_each_recommendation():
    recs = build_recommendations([_fail("nulls:age", 11, summary="5.4% null | with a pipe")], n_rows=204)
    md = render_markdown_actions(recs, "sample")
    assert "# Recommended Actions — sample" in md
    assert "| age | Missing values | 11 (5.4%) |" in md
    assert "with a pipe" in md and "\\|" in md  # pipes escaped so the table doesn't break
    assert "Nothing in the uploaded data has been changed" in md


def test_markdown_when_nothing_found():
    assert "No data quality issues were found" in render_markdown_actions([], "sample")


# --- failed validation rules become recommended actions --------------------

from dq_framework.expectations import ValidationOutcome


def _rule(exp_type: str, column: str | None, n: int, pct: float = 0.1, params: dict | None = None, success: bool = False):
    return ValidationOutcome(exp_type, column, success, n, pct, "summary", params or {})


def test_failed_unique_rule_becomes_a_recommendation():
    recs = build_recommendations(
        [], n_rows=1020, validation_outcomes=[_rule("expect_column_values_to_be_unique", "Email", 1020, 1.0)]
    )
    (rec,) = recs
    assert rec.column == "Email"
    assert rec.issue == "Values that should be unique repeat"
    assert rec.affected_rows == 1020
    assert rec.pct_of_rows == 1.0
    assert "unique" in rec.finding


def test_each_rule_type_maps_to_an_issue():
    outcomes = [
        _rule("expect_column_values_to_be_in_set", "status", 4, params={"value_set": ["a", "b"]}),
        _rule("expect_column_values_to_be_between", "age", 5, params={"min_value": 0, "max_value": 120}),
        _rule("expect_column_values_to_match_regex", "zip", 6, params={"regex": r"^\d{5}$"}),
        _rule("expect_table_row_count_to_be_between", None, 0, params={"min_value": 500, "max_value": 900}),
    ]
    recs = {r.issue: r for r in build_recommendations([], n_rows=100, validation_outcomes=outcomes)}
    assert set(recs) == {
        "Values outside the allowed set",
        "Values outside the expected range",
        "Values not matching the expected pattern",
        "Row count outside the expected range",
    }
    assert "'a', 'b'" in recs["Values outside the allowed set"].finding
    assert "[0, 120]" in recs["Values outside the expected range"].finding
    assert recs["Row count outside the expected range"].column is None


def test_passing_rules_produce_nothing():
    ok = _rule("expect_column_values_to_be_unique", "id", 0, success=True)
    assert build_recommendations([], n_rows=100, validation_outcomes=[ok]) == []


def test_rule_failure_already_reported_by_an_anomaly_check_is_not_repeated():
    anomalies = [
        _fail("nulls:age", 211),
        _fail("outliers:salary", 7),
        _fail("duplicates:key[id]", 16),
    ]
    outcomes = [
        _rule("expect_column_values_to_not_be_null", "age", 211),
        _rule("expect_column_values_to_be_between", "salary", 7),
        _rule("expect_column_values_to_be_unique", "id", 16),
    ]
    recs = build_recommendations(anomalies, n_rows=1000, validation_outcomes=outcomes)
    assert sorted(r.issue for r in recs) == ["Duplicate keys", "Missing values", "Outliers"]


def test_custom_rule_on_same_column_with_different_findings_is_kept():
    """A hand-written range rule that flags different rows than the outlier
    check is its own finding, not a repeat."""
    recs = build_recommendations(
        [_fail("outliers:salary", 7)],
        n_rows=1000,
        validation_outcomes=[
            _rule("expect_column_values_to_be_between", "salary", 40, params={"min_value": 0, "max_value": 100})
        ],
    )
    assert sorted(r.issue for r in recs) == ["Outliers", "Values outside the expected range"]


def test_reports_show_every_recommendation_and_label_custom_rules():
    """Every recommended action reaches all three reports, and the rule
    tables say which rules were custom and what they required."""
    import io

    import openpyxl

    from dq_framework import pipeline
    from dq_framework.profiling import DatasetProfile

    anomalies = [_fail("nulls:age", 11), _fail("duplicates:exact_rows", 8)]
    outcomes = [
        _rule("expect_column_values_to_be_in_set", "status", 4, params={"value_set": ["a", "b"]}),
        _rule("expect_column_values_to_be_unique", "email", 9),
    ]
    outcomes[0].source = "custom"
    recs = build_recommendations(anomalies, n_rows=100, validation_outcomes=outcomes)
    assert len(recs) == 4

    profile = DatasetProfile(n_rows=100, n_columns=0, columns={})
    reports = pipeline.generate_reports("t", anomalies, outcomes, recs, profile)

    assert reports["recommended_actions.md"].decode().count("\n| ") - 1 == len(recs)
    wb = openpyxl.load_workbook(io.BytesIO(reports["dq_scorecard.xlsx"]))
    assert wb["Recommended Actions"].max_row - 1 == len(recs)
    types = [row[1] for row in wb["Rule Summary"].iter_rows(min_row=2, values_only=True)]
    assert "validation rule (custom)" in types and "validation rule (auto-generated)" in types
    html = reports["dq_report.html"].decode()
    assert "one of &#x27;a&#x27;, &#x27;b&#x27;" in html and "custom" in html
